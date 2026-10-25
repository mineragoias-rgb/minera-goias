# -*- coding: utf-8 -*-
"""Checagens sem rede.

`checar_repositorio`: os arquivos do P3M e da ANM que estão no repositório seguem o formato que o site e os geradores usam?
  - P3M: cada recortado existe, tem .shp/.shx/.dbf/.prj, está em SIRGAS 2000 com .cpg UTF-8, tem o número de registros do
    `_indice_recorte.csv` e traz os campos que `scripts/mapa_camadas.json` lê; a chave de cada camada do mapa é única e
    preenchida no conjunto dos zips dela; cada camada do mapa tem URL e data no inventário.
  - ANM: o recorte de Goiás refeito sobre o próprio arquivo não muda nenhum byte (CFEM só com UF = GO; Cadastro Mineiro com as
    contagens de `recorte_goias.json`); os CSV abrem com a codificação e o separador esperados; o SIGMINE tem PROCESSO e está em
    SIRGAS 2000; cada arquivo tem data em `datas_de_acesso.json` e o sha256 e o tamanho certos no LEIA-ME de `Dados brutos`.
`conferir_planilha`: a base consolidada refeita pelo pipeline ganhou algum ALERTA novo na aba 14b?
"""
import csv
import json
import re
from pathlib import Path

import anm
import anm_recorte
import forma
import p3m as p3m_robo
from comum import sha256

P3M_REL = Path("Squad 1") / "P3M"
BRUTOS_REL = Path("Squad 1") / "Dados brutos"
CONFIG = Path(__file__).resolve().parent / "fontes.json"


def campos_por_camada(raiz):
    """{id da camada do mapa: (zips, campos da fonte que ela lê)}."""
    cfg = json.loads((Path(raiz) / "scripts" / "mapa_camadas.json").read_text(encoding="utf-8"))
    saida = {}
    for camada in cfg["camadas"]:
        origem = camada.get("origem") or {}
        if origem.get("tipo") != "p3m":
            continue
        campos = set()
        for c in camada.get("campos", []):
            f = c.get("fonte")
            campos |= set(f) if isinstance(f, list) else ({f} if f else set())
        campos |= {v for v in (camada.get("titular") or {}).values() if isinstance(v, str)}
        saida[camada["id"]] = ([a["arquivo"] for a in origem.get("arquivos", [])], campos)
    return saida


def checar_p3m(raiz):
    problemas = []
    p3m = Path(raiz) / P3M_REL
    rec_dir = p3m / "Recortado pelo polígono de Goiás"
    with open(rec_dir / "_indice_recorte.csv", encoding="utf-8-sig", newline="") as f:
        indice = {r["arquivo"]: r for r in csv.DictReader(f, delimiter=";")}
    lidos = {}
    for nome, r in indice.items():
        achados = [z for z in rec_dir.glob("*/*.zip") if z.name == nome]
        if len(achados) != 1:
            problemas.append(f"P3M: {nome} aparece {len(achados)} vezes no Recortado")
            continue
        try:
            partes = forma.ler_shapefile_zip(achados[0])
        except Exception as e:  # noqa: BLE001 — qualquer falha de leitura é problema de forma
            problemas.append(f"P3M: {nome} não abre: {e}")
            continue
        lidos[nome] = partes
        n = sum(p["n"] for p in partes)
        if n != int(r["registros_em_goias"]):
            problemas.append(f"P3M: {nome} tem {n} registros; o índice diz {r['registros_em_goias']}")
        for p in partes:
            if p["faltam"]:
                problemas.append(f"P3M: {nome} sem {', '.join(p['faltam'])}")
            if forma.crs_de(p["prj"]) != "SIRGAS 2000":
                problemas.append(f"P3M: {nome} em {forma.crs_de(p['prj'])}, não em SIRGAS 2000")
            if p["codificacao"].upper().replace("-", "") != "UTF8":
                problemas.append(f"P3M: {nome} sem .cpg UTF-8 ({p['codificacao'] or 'nenhum'})")
    for camada, (zips, campos) in campos_por_camada(raiz).items():
        presentes = set()
        for z in zips:
            if z not in lidos:
                problemas.append(f"P3M: a camada {camada} do mapa lê {z}, que não está no Recortado")
                continue
            presentes |= {c[0] for p in lidos[z] for c in p["campos"]}
        faltam = sorted(campos - presentes)
        if faltam and all(z in lidos for z in zips):
            problemas.append(f"P3M: a camada {camada} do mapa lê campos que os zips não têm: {faltam}")
    for camada, (zips, campo) in p3m_robo.chaves_do_site(raiz).items():
        valores = [v for z in zips for c in rec_dir.glob(f"*/{p3m_robo.glob_escape(z)}") for v in p3m_robo.valores_do_campo(c, campo)]
        problema = p3m_robo.problema_de_chave(campo, valores)
        if problema:
            problemas.append(f"P3M: a camada {camada} do mapa tem {problema}")
    with open(p3m / "inventario_p3m.csv", encoding="utf-8-sig", newline="") as f:
        inventario = list(csv.DictReader(f, delimiter=";"))
    no_mapa = {z for zs, _ in campos_por_camada(raiz).values() for z in zs}
    for ln in inventario:
        if {Path(a).name for a in ln["arquivo_recortado"].split(";") if a} & no_mapa:
            if not ln["url_download"]:
                problemas.append(f"P3M: {ln['camada']} está no mapa e não tem URL de download no inventário")
            if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", ln["data_download"]):
                problemas.append(f"P3M: {ln['camada']} sem data de download no inventário")
    return problemas


def checar_anm(raiz):
    problemas = []
    raiz = Path(raiz)
    brutos = raiz / BRUTOS_REL
    cfg = json.loads(CONFIG.read_text(encoding="utf-8"))["anm"]
    datas = json.loads((brutos / "datas_de_acesso.json").read_text(encoding="utf-8"))
    leiame = (brutos / "LEIA-ME.md").read_text(encoding="utf-8")
    recorte = json.loads((brutos / "ANM - Cadastro Mineiro" / "recorte_goias.json").read_text(encoding="utf-8"))
    arquivos = []  # (caminho, lógico)
    for f in cfg["fontes"]:
        if f["tipo"] == "shapefile":
            caminho = brutos / f["arquivo"]
            partes = forma.ler_shapefile_zip(caminho)
            campos = {c[0] for p in partes for c in p["campos"]}
            if "PROCESSO" not in campos:
                problemas.append(f"ANM: {f['arquivo']} sem o campo PROCESSO")
            if {forma.crs_de(p["prj"]) for p in partes} != {"SIRGAS 2000"}:
                problemas.append(f"ANM: {f['arquivo']} fora de SIRGAS 2000")
            arquivos.append((caminho, f["logico"]))
        elif f["tipo"] == "cadastro":
            municipios = anm_recorte.municipios_goias(raiz / f["malha_municipal"])
            processos = anm_recorte.processos_sigmine(brutos / f["sigmine"])
            for item in recorte["arquivos"]:
                caminho = brutos / f["pasta"] / item["arquivo"]
                dados = caminho.read_bytes()
                novo, cont = anm_recorte.recortar_cadastro(dados, item["regra"], municipios, processos, f["separador"])
                if novo != dados:
                    problemas.append(f"ANM: {item['arquivo']} tem linhas fora da regra de recorte de Goiás")
                for k, v in cont.items():
                    if k != "linhas_nacional" and item.get(k) != v:
                        problemas.append(f"ANM: {item['arquivo']}: {k} = {v}, recorte_goias.json diz {item.get(k)}")
                if item.get("sha256_goias") != sha256(caminho):
                    problemas.append(f"ANM: {item['arquivo']}: sha256 diferente do registrado em recorte_goias.json")
                problemas += [f"ANM: {item['arquivo']}: {b}" for b in forma.checar_csv(dados, dados, f["codificacao"], f["separador"])[0]]
                arquivos.append((caminho, f["logico"] + item["arquivo"]))
        else:
            caminho = raiz / f["arquivo_raiz"] if f.get("arquivo_raiz") else brutos / f["arquivo"]
            dados = caminho.read_bytes()
            if f.get("recorte"):
                novo, _ = anm_recorte.recortar_uf(dados, f["separador"], f["recorte"]["coluna"], f["recorte"]["valor"])
                if novo != dados:
                    problemas.append(f"ANM: {f['arquivo']} tem linhas fora de {f['recorte']['coluna']} = {f['recorte']['valor']}")
            problemas += [f"ANM: {caminho.name}: {b}" for b in forma.checar_csv(dados, dados, f["codificacao"], f["separador"])[0]]
            if f.get("periodo") and periodo_no_leiame(leiame, f["src"]) != anm.periodo_csv(dados, f):
                problemas.append(f"ANM: período de {f['src']} no LEIA-ME de Dados brutos ({periodo_no_leiame(leiame, f['src'])}) "
                                 f"difere do arquivo ({anm.periodo_csv(dados, f)})")
            if f.get("logico"):  # as rodadas (Squad 1/dados) não têm data nem linha no LEIA-ME de Dados brutos
                arquivos.append((caminho, f["logico"]))
    return problemas + registro_em_dia(arquivos, datas, leiame, "ANM")


def registro_em_dia(arquivos, datas, leiame, rotulo):
    """Cada arquivo tem data em datas_de_acesso.json e o tamanho e o sha256 certos no LEIA-ME de Dados brutos."""
    problemas = []
    for caminho, logico in arquivos:
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", datas.get(logico, "")):
            problemas.append(f"{rotulo}: {logico} sem data em datas_de_acesso.json")
        mb = f"{caminho.stat().st_size / 1e6:.1f}".replace(".", ",")
        if f"| `{caminho.name}` | {mb} MB | `{sha256(caminho)}` |" not in leiame:
            problemas.append(f"{rotulo}: {caminho.name}: tamanho ou sha256 desatualizado no LEIA-ME de Dados brutos")
    return problemas


def periodo_no_leiame(leiame, src):
    """A célula 'Período coberto' da linha da fonte na tabela de fontes do LEIA-ME."""
    for ln in leiame.split("\n"):
        if f"**{src}**" in ln and ln.startswith("| `"):
            return ln.split(" | ")[3]
    return None


def checar_diretos(raiz):
    """SIDRA, RECMIN e malhas: forma do arquivo, período no LEIA-ME, metadados do RECMIN, datas e sha256."""
    import diretos
    problemas, arquivos = [], []
    raiz = Path(raiz)
    brutos = raiz / BRUTOS_REL
    cfg = json.loads(CONFIG.read_text(encoding="utf-8"))["diretos"]
    datas = json.loads((brutos / "datas_de_acesso.json").read_text(encoding="utf-8"))
    leiame = (brutos / "LEIA-ME.md").read_text(encoding="utf-8")
    for f in cfg["fontes"]:
        rotulo = f["grupo"].upper()
        if f["tipo"] == "malha":
            if not (raiz / f["arquivo_raiz"]).exists():
                problemas.append(f"{rotulo}: {f['arquivo_raiz']} não existe")
            continue
        caminho = brutos / f["arquivo"]
        dados = caminho.read_bytes()
        if f["tipo"] == "sidra":
            bloq, _, _ = diretos.checar_sidra(dados, dados)
            periodo = diretos.periodo_sidra(json.loads(dados), f.get("periodo", "municipal"))
        else:
            bloq, _ = diretos.checar_geojson(dados, dados)
            meta_path = caminho.with_name(caminho.name.replace(".geojson", ".metadados.json"))
            meta = json.loads(meta_path.read_text(encoding="utf-8"))
            j = json.loads(dados)
            for campo, valor in (("sha256", sha256(caminho)), ("bytes", len(dados)), ("feicoes", len(j["features"]))):
                if meta.get(campo) != valor:
                    problemas.append(f"{rotulo}: {meta_path.name}: {campo} = {meta.get(campo)}, o arquivo tem {valor}")
            periodo = diretos.periodo_recmin(j, meta["data_download_utc"])
            arquivos.append((meta_path, f["logico_metadados"]))
        problemas += [f"{rotulo}: {caminho.name}: {b}" for b in bloq]
        if periodo_no_leiame(leiame, f["src"]) != periodo:
            problemas.append(f"{rotulo}: período de {f['src']} no LEIA-ME de Dados brutos ({periodo_no_leiame(leiame, f['src'])}) "
                             f"difere do arquivo ({periodo})")
        arquivos.append((caminho, f["logico"]))
    return problemas + registro_em_dia(arquivos, datas, leiame, "IBGE/SGB")


def checar_repositorio(raiz):
    return checar_p3m(raiz) + checar_anm(raiz) + checar_diretos(raiz)


# ---------------------------------------------------------------------------------------------------------- planilha
def alertas_14b(caminho):
    import openpyxl
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    linhas = list(wb["14b_validacoes_governanca"].iter_rows(values_only=True))
    cab = next(i for i, ln in enumerate(linhas) if ln and ln[0] == "verificacao")
    return {str(ln[0]): (str(ln[3]), str(ln[2])) for ln in linhas[cab + 1:] if ln and ln[0]}


def conferir_planilha(antes, depois):
    """Bloqueia quando a planilha refeita tem ALERTA numa verificação que não tinha (ou numa verificação nova).

    `planilha.alertas_esperados` em robo/fontes.json lista os que valem no arranjo do GitHub, com o motivo."""
    esperados = json.loads(CONFIG.read_text(encoding="utf-8")).get("planilha", {}).get("alertas_esperados", {})
    a, d = alertas_14b(antes), alertas_14b(depois)
    problemas, linhas = [], ["| Verificação | Antes | Depois |", "|---|---|---|"]
    for nome, (status, resultado) in d.items():
        anterior = a.get(nome, ("(nova)", ""))[0]
        if status != "OK" or anterior != "OK":
            linhas.append(f"| {nome[:90]} | {anterior} | {status} |")
        if status != "OK" and anterior != status and not any(nome.startswith(e) for e in esperados):
            problemas.append(f"{nome}: {anterior} → {status} ({resultado[:200]})")
    for nome in a:
        if nome not in d:
            problemas.append(f"verificação sumiu da aba 14b: {nome}")
    texto = "\n".join(linhas) + "\n\n" + ("\n".join("BLOQUEIO: " + p for p in problemas) or "nenhum alerta novo na aba 14b")
    return problemas, texto


def mesmo_conteudo_xlsx(a, b):
    """True quando as duas planilhas têm as mesmas abas e os mesmos valores (o .xlsx muda de bytes a cada gravação, pela data)."""
    import openpyxl
    wa = openpyxl.load_workbook(a, read_only=True, data_only=True)
    wb = openpyxl.load_workbook(b, read_only=True, data_only=True)
    if wa.sheetnames != wb.sheetnames:
        return False
    return all(list(wa[n].iter_rows(values_only=True)) == list(wb[n].iter_rows(values_only=True)) for n in wa.sheetnames)

