#!/usr/bin/env python3
"""Incremental source-data inventory. Does not execute repository code or SQL."""
import argparse
import csv
from datetime import date, datetime, timezone
from decimal import Decimal
import fnmatch
import hashlib
import io
import itertools
import json
import math
import os
from pathlib import Path
import re
import sqlite3
import sys
import unicodedata
import uuid
import zipfile

PARSER_VERSION = '1'
SUPPORTED = {'.csv', '.tsv', '.xlsx'}
EXCLUDED_DIRS = {'.git', '.venv', 'venv', 'node_modules', '__pycache__', '.cache', '.pytest_cache'}


def encode(value):
    def default(obj):
        if isinstance(obj, (datetime, date)):
            return obj.isoformat()
        if isinstance(obj, Decimal):
            return str(obj)
        raise TypeError(type(obj).__name__)
    return json.dumps(value, ensure_ascii=False, default=default, allow_nan=False, separators=(',', ':'))


def digest(*parts):
    return hashlib.sha256(encode(parts).encode()).hexdigest()


def timestamp():
    return datetime.now(timezone.utc).isoformat()


def schema(mysql=False):
    long = 'LONGTEXT' if mysql else 'TEXT'
    tail = ' ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_bin' if mysql else ''
    tables = [
        f'''CREATE TABLE IF NOT EXISTS ingest_runs (
            run_id CHAR(32) PRIMARY KEY, repo_id VARCHAR(255) NOT NULL,
            commit_sha VARCHAR(64) NOT NULL, started_at VARCHAR(40) NOT NULL,
            finished_at VARCHAR(40), status VARCHAR(32) NOT NULL, summary_json {long})''',
        f'''CREATE TABLE IF NOT EXISTS ingest_files (
            file_id CHAR(64) PRIMARY KEY, repo_id VARCHAR(255) NOT NULL,
            path {long} NOT NULL, extension VARCHAR(20), content_sha CHAR(64),
            size_bytes BIGINT, classification VARCHAR(32), status VARCHAR(32),
            active INTEGER NOT NULL, current_version CHAR(64), last_seen_run CHAR(32),
            last_error {long})''',
        f'''CREATE TABLE IF NOT EXISTS ingest_versions (
            version_id CHAR(64) PRIMARY KEY, file_id CHAR(64) NOT NULL,
            content_sha CHAR(64) NOT NULL, commit_sha VARCHAR(64) NOT NULL,
            parser_version VARCHAR(20) NOT NULL, policy_sha CHAR(64) NOT NULL,
            imported_at VARCHAR(40) NOT NULL, dataset_count INTEGER, row_count BIGINT)''',
        f'''CREATE TABLE IF NOT EXISTS ingest_datasets (
            dataset_id CHAR(64) PRIMARY KEY, version_id CHAR(64) NOT NULL,
            sheet_name VARCHAR(255) NOT NULL, header_row INTEGER NOT NULL,
            columns_json {long} NOT NULL, metadata_json {long} NOT NULL,
            row_count BIGINT, data_row_count BIGINT, warnings_json {long})''',
        f'''CREATE TABLE IF NOT EXISTS ingest_rows (
            dataset_id CHAR(64) NOT NULL, source_row INTEGER NOT NULL,
            row_role VARCHAR(20) NOT NULL, values_json {long} NOT NULL,
            formulas_json {long} NOT NULL,
            PRIMARY KEY(dataset_id, source_row))'''
    ]
    view = '''CREATE VIEW IF NOT EXISTS ingest_current_rows AS
        SELECT f.repo_id, f.path, f.classification, f.status AS import_status,
               v.commit_sha, v.content_sha, d.dataset_id, d.sheet_name,
               d.columns_json, r.source_row, r.row_role, r.values_json, r.formulas_json
        FROM ingest_files f JOIN ingest_versions v ON f.current_version=v.version_id
        JOIN ingest_datasets d ON d.version_id=v.version_id
        JOIN ingest_rows r ON r.dataset_id=d.dataset_id WHERE f.active=1'''
    if mysql:
        view = view.replace('CREATE VIEW IF NOT EXISTS', 'CREATE OR REPLACE VIEW')
    return [s + tail for s in tables] + [view]


class Database:
    def __init__(self, connection, mysql=False):
        self.connection, self.mysql = connection, mysql

    def execute(self, sql, args=()):
        cursor = self.connection.cursor()
        cursor.execute(sql.replace('?', '%s') if self.mysql else sql, args)
        return cursor

    def one(self, sql, args=()):
        cursor = self.execute(sql, args)
        try:
            return cursor.fetchone()
        finally:
            cursor.close()

    def many(self, sql, rows):
        cursor = self.connection.cursor()
        try:
            cursor.executemany(sql.replace('?', '%s') if self.mysql else sql, rows)
        finally:
            cursor.close()

    def write(self, sql, args=()):
        self.execute(sql, args).close()

    def commit(self):
        self.connection.commit()

    def rollback(self):
        self.connection.rollback()

    def initialize(self):
        for statement in schema(self.mysql):
            self.write(statement)
        self.commit()


def discover(root):
    paths = []
    def onerror(error):
        raise error
    for base, dirs, files in os.walk(root, followlinks=False, onerror=onerror):
        dirs[:] = sorted(d for d in dirs if d not in EXCLUDED_DIRS and not d.startswith('.')
                         and not (Path(base) / d).is_symlink())
        for name in sorted(files):
            path = Path(base) / name
            if name.startswith('.') or name.startswith('~$') or path.is_symlink():
                continue
            paths.append(path)
    return paths


def classification(path):
    if re.search(r'(^|[/_ .-])(demo|mock|sintetico|synthetic)([/_ .-]|$)', str(path).lower()):
        return 'demo'
    return 'source_unvalidated' if path.suffix.lower() in SUPPORTED else 'document_or_code'


def csv_rows(path):
    raw = path.read_bytes()
    encoding = None
    for candidate in ('utf-8-sig', 'cp1252', 'latin1'):
        try:
            text = raw.decode(candidate)
            encoding = candidate
            break
        except UnicodeDecodeError:
            continue
    if '\x00' in text:
        raise ValueError('NUL bytes: unsupported text encoding')
    first = next((line for line in text.splitlines() if line.strip()), '')
    delimiter = '\t' if path.suffix.lower() == '.tsv' else max((';', ',', '\t'), key=first.count)
    reader = csv.reader(io.StringIO(text, newline=''), delimiter=delimiter, strict=True)
    def rows():
        # source_row is the starting physical line, including multiline fields.
        line = 1
        for values in reader:
            yield line, values, {}
            line = reader.line_num + 1
    return rows(), {'encoding': encoding, 'delimiter': delimiter}


def workbook_sheets(path, policy):
    from openpyxl import load_workbook
    with zipfile.ZipFile(path) as archive:
        if sum(i.file_size for i in archive.infolist()) > policy['max_uncompressed_bytes']:
            raise ValueError('Workbook exceeds uncompressed size limit')
    values = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    formulas = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        for sheet in values:
            other = formulas[sheet.title]
            def rows(sheet=sheet, other=other):
                for number, (vr, fr) in enumerate(zip(sheet.iter_rows(), other.iter_rows()), 1):
                    vv = [cell.value for cell in vr]
                    ff = {i: cell.value if isinstance(cell.value, str) else str(cell.value)
                          for i, cell in enumerate(fr) if cell.data_type == 'f'}
                    yield number, vv, ff
            yield sheet.title, rows(), {'format': 'xlsx', 'formula_policy': 'cached values; formulas preserved, never executed'}
    finally:
        values.close()
        formulas.close()


def header_for(path, sheet, sample, policy):
    for rule in policy['rules']:
        if fnmatch.fnmatch(path.name, rule['file']) and fnmatch.fnmatch(sheet, rule['sheet']):
            return rule['header_row'], 'configured'
    if path.suffix.lower() in {'.csv', '.tsv'}:
        return next((n for n, row, _ in sample if any(x not in (None, '') for x in row)), 0), 'first_nonempty_row'
    for n, row, _ in sample:
        populated = [v for v in row if v not in (None, '')]
        if len(populated) >= 2 and all(isinstance(v, str) for v in populated):
            return n, 'inferred_requires_review'
    return 0, 'no_header_detected'


def columns_for(values, width):
    result, used = [], set()
    for index in range(width):
        label = values[index] if index < len(values) else None
        base = unicodedata.normalize('NFKD', str(label or ''))
        base = re.sub('[^a-z0-9]+', '_', base.encode('ascii', 'ignore').decode().lower()).strip('_')
        base = base or f'col_{index + 1}'
        key, suffix = base, 2
        while key in used:
            key = f'{base}_{suffix}'
            suffix += 1
        used.add(key)
        result.append({'key': key, 'label': label, 'position': index + 1})
    return result


def import_sheet(db, path, version, name, rows, metadata, policy):
    iterator = iter(rows)
    sample = list(itertools.islice(iterator, 50))
    header, method = header_for(path, name, sample, policy)
    labels = next((v for n, v, _ in sample if n == header), [])
    width = max((len(v) for _, v, _ in sample), default=0)
    columns = columns_for(labels, width)
    warnings = {'missing_formula_cache': 0, 'scientific_identifier': 0, 'numeric_identifier': 0,
                'spreadsheet_errors': 0, 'extra_columns': 0}
    dataset = digest(version, name)
    metadata.update({'header_method': method, 'validation': 'not_validated',
                     'numeric_policy': 'CSV text and Excel stored values preserved; no inferred conversions'})
    db.write('INSERT INTO ingest_datasets (dataset_id,version_id,sheet_name,header_row,columns_json,metadata_json) VALUES (?,?,?,?,?,?)',
             (dataset, version, name, header, encode(columns), encode(metadata)))
    batch, count, data_count = [], 0, 0
    for number, values, formulas in itertools.chain(sample, iterator):
        if number > policy['max_rows_per_sheet']:
            raise ValueError('Sheet exceeds row limit')
        if not any(v not in (None, '') for v in values) and not formulas:
            continue
        if len(values) > len(columns):
            # Keep every column; never truncate data after an uneven header.
            extra = len(values) - len(columns)
            columns = columns_for(labels, len(values))
            warnings['extra_columns'] += extra
        role = 'context' if not header or number < header else ('header' if number == header else 'data')
        payload, stored_formulas = {}, {}
        for i, col in enumerate(columns):
            value = values[i] if i < len(values) else None
            if isinstance(value, float) and not math.isfinite(value):
                raise ValueError('Non-finite numeric cell')
            payload[col['key']] = value
            if i in formulas:
                stored_formulas[col['key']] = formulas[i]
                if value is None:
                    warnings['missing_formula_cache'] += 1
            if role == 'data' and any(token in col['key'] for token in ('cnpj', 'cpf', 'documento')):
                if isinstance(value, str) and re.search(r'\d[.,]?\d*[eE][+-]?\d+', value):
                    warnings['scientific_identifier'] += 1
                if isinstance(value, (int, float)):
                    warnings['numeric_identifier'] += 1
            if isinstance(value, str) and value in {'#REF!', '#DIV/0!', '#VALUE!', '#N/A', '#NAME?', '#NUM!', '#NULL!'}:
                warnings['spreadsheet_errors'] += 1
        batch.append((dataset, number, role, encode(payload), encode(stored_formulas)))
        count += 1
        data_count += role == 'data'
        if len(batch) >= 500:
            db.many('INSERT INTO ingest_rows VALUES (?,?,?,?,?)', batch)
            batch.clear()
    if batch:
        db.many('INSERT INTO ingest_rows VALUES (?,?,?,?,?)', batch)
    db.write('UPDATE ingest_datasets SET columns_json=?,row_count=?,data_row_count=?,warnings_json=? WHERE dataset_id=?',
             (encode(columns), count, data_count, encode(warnings), dataset))
    return {'sheet': name, 'rows': count, 'data_rows': data_count, 'header_row': header, 'warnings': warnings}


def scan(db, root, repo_id, commit, policy):
    root = Path(root).resolve(strict=True)
    files = discover(root)  # Discovery must succeed before deactivating removed files.
    run = uuid.uuid4().hex
    summary = {'repo_id': repo_id, 'commit': commit, 'run_id': run, 'files': [], 'errors': 0,
               'imported_files': 0, 'unchanged_files': 0, 'imported_rows': 0}
    policy_sha = digest(policy)
    db.write('INSERT INTO ingest_runs (run_id,repo_id,commit_sha,started_at,status) VALUES (?,?,?,?,?)',
             (run, repo_id, commit, timestamp(), 'running'))
    db.commit()
    for path in files:
        relative = path.relative_to(root).as_posix()
        file_id = digest(repo_id, relative)
        kind = classification(Path(relative))
        entry = {'path': relative, 'classification': kind}
        try:
            size = path.stat().st_size
            if size > policy['max_file_bytes']:
                raise ValueError('File exceeds size limit')
            h = hashlib.sha256()
            with path.open('rb') as source:
                for chunk in iter(lambda: source.read(1024 * 1024), b''):
                    h.update(chunk)
            sha = h.hexdigest()
            previous = db.one('SELECT current_version FROM ingest_files WHERE file_id=?', (file_id,))
            if previous is None:
                db.write('INSERT INTO ingest_files (file_id,repo_id,path,active) VALUES (?,?,?,1)', (file_id, repo_id, relative))
            db.write('UPDATE ingest_files SET extension=?,content_sha=?,size_bytes=?,classification=?,active=1,last_seen_run=?,last_error=NULL WHERE file_id=?',
                     (path.suffix.lower(), sha, size, kind, run, file_id))
            if path.suffix.lower() not in SUPPORTED or (kind == 'demo' and policy['exclude_demo']):
                state = 'demo_excluded' if kind == 'demo' else 'catalog_only'
                db.write('UPDATE ingest_files SET status=?,current_version=NULL WHERE file_id=?', (state, file_id))
                entry['status'] = state
            else:
                version = digest(file_id, sha, PARSER_VERSION, policy_sha)
                exists = db.one('SELECT version_id FROM ingest_versions WHERE version_id=?', (version,))
                if exists:
                    entry['status'] = 'unchanged'
                    summary['unchanged_files'] += 1
                else:
                    db.write('INSERT INTO ingest_versions (version_id,file_id,content_sha,commit_sha,parser_version,policy_sha,imported_at) VALUES (?,?,?,?,?,?,?)',
                             (version, file_id, sha, commit, PARSER_VERSION, policy_sha, timestamp()))
                    if path.suffix.lower() == '.xlsx':
                        sheets = workbook_sheets(path, policy)
                    else:
                        rows, meta = csv_rows(path)
                        sheets = iter([('data', rows, meta)])
                    stats = []
                    try:
                        for name, rows, meta in sheets:
                            stats.append(import_sheet(db, path, version, name, rows, meta, policy))
                    finally:
                        if hasattr(sheets, 'close'):
                            sheets.close()
                    total = sum(s['rows'] for s in stats)
                    db.write('UPDATE ingest_versions SET dataset_count=?,row_count=? WHERE version_id=?', (len(stats), total, version))
                    summary['imported_rows'] += total
                    summary['imported_files'] += 1
                    entry.update(status='imported', datasets=stats)
                db.write('UPDATE ingest_files SET current_version=?,status=? WHERE file_id=?', (version, 'imported', file_id))
            db.commit()  # Publish each file only when all of its sheets succeed.
        except Exception as error:
            db.rollback()
            summary['errors'] += 1
            message = f'{type(error).__name__}: {error}'[:1000]
            entry.update(status='error', error=message)
            if not db.one('SELECT file_id FROM ingest_files WHERE file_id=?', (file_id,)):
                db.write('INSERT INTO ingest_files (file_id,repo_id,path,active,classification) VALUES (?,?,?,1,?)',
                         (file_id, repo_id, relative, kind))
            db.write('UPDATE ingest_files SET active=1,status=?,last_error=?,last_seen_run=? WHERE file_id=?',
                     ('error', message, run, file_id))
            db.commit()
        summary['files'].append(entry)
    db.write('UPDATE ingest_files SET active=0,status=? WHERE repo_id=? AND (last_seen_run<>? OR last_seen_run IS NULL)',
             ('removed', repo_id, run))
    db.write('UPDATE ingest_runs SET finished_at=?,status=?,summary_json=? WHERE run_id=?',
             (timestamp(), 'partial' if summary['errors'] else 'success', encode(summary), run))
    db.commit()
    return summary


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--root', type=Path)
    parser.add_argument('--repo-id', default='mineragoias-rgb/minera-goias')
    parser.add_argument('--commit', default='working-tree')
    parser.add_argument('--policy', type=Path, default=Path(__file__).with_name('policy.json'))
    parser.add_argument('--sqlite', type=Path)
    parser.add_argument('--mysql', action='store_true')
    parser.add_argument('--schema-mysql', action='store_true')
    parser.add_argument('--report', type=Path)
    args = parser.parse_args()
    if args.schema_mysql:
        print(';\n'.join(schema(True)) + ';')
        return 0
    if args.root is None or bool(args.sqlite) == args.mysql:
        parser.error('Use --root and exactly one of --sqlite / --mysql')
    if args.report and args.report.resolve().is_relative_to(args.root.resolve()):
        parser.error('Report must be outside the inspected repository')
    if args.sqlite:
        if args.sqlite.resolve().is_relative_to(args.root.resolve()):
            parser.error('Database must be outside the inspected repository')
        args.sqlite.parent.mkdir(parents=True, exist_ok=True)
        conn = sqlite3.connect(args.sqlite)
        db = Database(conn)
        db.initialize()
    else:
        import pymysql
        conn = pymysql.connect(host=os.environ.get('DB_HOST', 'localhost'),
                               user=os.environ['DB_USER'], password=os.environ['DB_PASSWORD'],
                               database=os.environ['DB_NAME'], charset='utf8mb4', autocommit=False)
        db = Database(conn, True)  # Schema must be installed separately by administrator.
    try:
        if args.mysql and db.one('SELECT GET_LOCK(?,0)', ('ingest-' + digest(args.repo_id)[:48],))[0] != 1:
            raise RuntimeError('Another import for this repository is running')
        summary = scan(db, args.root, args.repo_id, args.commit, json.loads(args.policy.read_text()))
    finally:
        conn.close()
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        temporary = args.report.with_suffix('.tmp')
        temporary.write_text(json.dumps(summary, ensure_ascii=False, indent=2))
        temporary.replace(args.report)
    print(encode({key: value for key, value in summary.items() if key != 'files'}))
    return 1 if summary['errors'] else 0


if __name__ == '__main__':
    raise SystemExit(main())
